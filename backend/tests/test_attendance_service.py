from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest

from app.models import AttendanceCode, AttendanceEntry, Holiday
from app.services import attendance_export_service, attendance_service, leave_service
from tests.conftest import grant_balance


@pytest.fixture()
def code_present(db) -> AttendanceCode:
    code = AttendanceCode(libelle="Présent", code_court="P", couleur="#43A047")
    db.add(code)
    db.flush()
    return code


def test_guess_identifier_column_matches_common_headers():
    assert attendance_service.guess_identifier_column(["Ville", "Nom", "1", "2"]) == "Nom"
    assert attendance_service.guess_identifier_column(["Matricule", "1"]) == "Matricule"
    assert attendance_service.guess_identifier_column(["Ville", "Poste"]) is None


def test_guess_day_columns_picks_1_to_31_only():
    columns = ["Nom", "01", "2", "31", "32", "Total"]
    assert attendance_service.guess_day_columns(columns) == ["01", "2", "31"]


def test_find_unmapped_day_values_flags_values_matching_no_code(db, code_present):
    # "PRST" doesn't match code_present's code_court ("P") or libelle
    # ("Présent") even after accent/case normalization — genuinely unmapped,
    # unlike e.g. "PRESENT" which normalizes to match "Présent".
    rows = [
        {"Nom": "Sara Alami", "01": "PRST", "02": "P"},
        {"Nom": "Karim Idrissi", "01": "P", "02": "XYZ"},
    ]
    unmapped = attendance_service.find_unmapped_day_values(db, rows, ["01", "02"])
    assert set(unmapped) == {"PRST", "XYZ"}


def test_find_unmapped_day_values_matches_by_libelle_too(db, code_present):
    rows = [{"Nom": "Sara Alami", "01": "Présent"}]
    unmapped = attendance_service.find_unmapped_day_values(db, rows, ["01"])
    assert unmapped == []


def test_find_unmapped_day_values_ignores_blank_cells(db, code_present):
    rows = [{"Nom": "Sara Alami", "01": "", "02": "  "}]
    unmapped = attendance_service.find_unmapped_day_values(db, rows, ["01", "02"])
    assert unmapped == []


def test_read_table_parses_csv():
    content = b"Nom,01,02\nSara Alami,P,A\n"
    columns, rows = attendance_service.read_table(content, "pointage.csv")
    assert columns == ["Nom", "01", "02"]
    assert rows == [{"Nom": "Sara Alami", "01": "P", "02": "A"}]


def test_read_table_rejects_unsupported_extension():
    with pytest.raises(attendance_service.AttendanceServiceError):
        attendance_service.read_table(b"whatever", "pointage.xls")


def test_run_import_matches_by_full_name_both_orders(
    db, employee_a, employee_b, hr_user, code_present
):
    rows = [
        {"Nom": f"{employee_a.prenom} {employee_a.nom}", "01": "P"},
        {"Nom": f"{employee_b.nom} {employee_b.prenom}", "01": "P"},
    ]
    import_row = attendance_service.run_import(
        db,
        rows=rows,
        identifier_column="Nom",
        day_columns=["01"],
        mois=9,
        annee=2026,
        filename="pointage.csv",
        actor_user_id=hr_user.id,
    )
    assert import_row.nb_lignes_importees == 2
    assert import_row.nb_lignes_non_reconnues == 0

    entry_a = (
        db.query(AttendanceEntry)
        .filter_by(employee_id=employee_a.id, date=date(2026, 9, 1))
        .first()
    )
    entry_b = (
        db.query(AttendanceEntry)
        .filter_by(employee_id=employee_b.id, date=date(2026, 9, 1))
        .first()
    )
    assert entry_a is not None and entry_a.code_id == code_present.id
    assert entry_b is not None and entry_b.code_id == code_present.id


def test_run_import_matches_by_matricule(db, employee_a, hr_user, code_present):
    employee_a.matricule = "MAT-042"
    db.flush()
    rows = [{"Matricule": "mat 042", "01": "P"}]
    import_row = attendance_service.run_import(
        db,
        rows=rows,
        identifier_column="Matricule",
        day_columns=["01"],
        mois=9,
        annee=2026,
        filename="pointage.csv",
        actor_user_id=hr_user.id,
    )
    assert import_row.nb_lignes_importees == 1
    entry = (
        db.query(AttendanceEntry)
        .filter_by(employee_id=employee_a.id, date=date(2026, 9, 1))
        .first()
    )
    assert entry is not None


def test_run_import_collects_unmatched_rows_without_writing(db, hr_user, code_present):
    rows = [{"Nom": "Personne Inconnue", "01": "P"}]
    import_row = attendance_service.run_import(
        db,
        rows=rows,
        identifier_column="Nom",
        day_columns=["01"],
        mois=9,
        annee=2026,
        filename="pointage.csv",
        actor_user_id=hr_user.id,
    )
    assert import_row.nb_lignes_importees == 0
    assert import_row.nb_lignes_non_reconnues == 1
    assert import_row.noms_non_reconnus == "Personne Inconnue"
    assert db.query(AttendanceEntry).count() == 0


def test_run_import_upserts_on_reimport_same_employee_date(db, employee_a, hr_user, code_present):
    rows = [{"Nom": f"{employee_a.prenom} {employee_a.nom}", "01": "P"}]
    attendance_service.run_import(
        db,
        rows=rows,
        identifier_column="Nom",
        day_columns=["01"],
        mois=9,
        annee=2026,
        filename="pointage1.csv",
        actor_user_id=hr_user.id,
    )
    absence = AttendanceCode(libelle="Absence", code_court="A", couleur="#E53935")
    db.add(absence)
    db.flush()
    rows2 = [{"Nom": f"{employee_a.prenom} {employee_a.nom}", "01": "A"}]
    attendance_service.run_import(
        db,
        rows=rows2,
        identifier_column="Nom",
        day_columns=["01"],
        mois=9,
        annee=2026,
        filename="pointage2.csv",
        actor_user_id=hr_user.id,
    )
    entries = (
        db.query(AttendanceEntry).filter_by(employee_id=employee_a.id, date=date(2026, 9, 1)).all()
    )
    assert len(entries) == 1
    assert entries[0].code_id == absence.id


def test_run_import_skips_invalid_calendar_day(db, employee_a, hr_user, code_present):
    # September has 30 days — day 31 must not raise, just be skipped.
    rows = [{"Nom": f"{employee_a.prenom} {employee_a.nom}", "31": "P"}]
    import_row = attendance_service.run_import(
        db,
        rows=rows,
        identifier_column="Nom",
        day_columns=["31"],
        mois=9,
        annee=2026,
        filename="pointage.csv",
        actor_user_id=hr_user.id,
    )
    assert import_row.nb_lignes_importees == 1
    assert db.query(AttendanceEntry).count() == 0


# ------------------------------------------------------------- export state --


def test_build_monthly_state_shows_pointage_code(db, employee_a, active_status, code_present):
    db.add(
        AttendanceEntry(employee_id=employee_a.id, date=date(2026, 9, 7), code_id=code_present.id)
    )
    db.flush()
    state = attendance_export_service.build_monthly_state(db, 9, 2026)
    row = next(r for r in state["rows"] if r["employee_id"] == employee_a.id)
    day7 = next(d for d in row["days"] if d["date"] == "2026-09-07")
    assert day7["code"] == "P"
    assert day7["conflict"] is False


def test_build_monthly_state_shows_holiday_code(db, employee_a):
    db.add(Holiday(date=date(2026, 9, 9), libelle="Test"))
    db.flush()
    state = attendance_export_service.build_monthly_state(db, 9, 2026)
    row = next(r for r in state["rows"] if r["employee_id"] == employee_a.id)
    day9 = next(d for d in row["days"] if d["date"] == "2026-09-09")
    assert day9["code"] == "F"


def test_build_monthly_state_flags_pointage_leave_conflict(
    db, employee_a, leave_type, hr_user, code_present
):
    grant_balance(db, employee_a.id, leave_type.id, 2026)
    request = leave_service.create_request(
        db,
        employee_id=employee_a.id,
        leave_type_id=leave_type.id,
        date_debut=date(2026, 9, 7),
        date_fin=date(2026, 9, 7),
        commentaire=None,
        submitted_by_user_id=hr_user.id,
    )
    leave_service.approve_request(db, request, decided_by_user_id=hr_user.id, comment=None)
    db.add(
        AttendanceEntry(employee_id=employee_a.id, date=date(2026, 9, 7), code_id=code_present.id)
    )
    db.flush()

    state = attendance_export_service.build_monthly_state(db, 9, 2026)
    row = next(r for r in state["rows"] if r["employee_id"] == employee_a.id)
    day7 = next(d for d in row["days"] if d["date"] == "2026-09-07")
    assert day7["conflict"] is True
    # The leave code (congé) takes precedence over the raw pointage code for
    # payroll purposes. The `leave_type` fixture has no code_court set, so
    # this also exercises the libelle-derived fallback abbreviation.
    assert day7["code"] == "CON"
    assert state["nb_conflits"] == 1


def test_export_monthly_state_xlsx_produces_a_real_workbook(db, employee_a):
    content = attendance_export_service.export_monthly_state_xlsx(db, 9, 2026)
    assert content[:2] == b"PK"  # xlsx is a zip archive
    assert len(content) > 1000


def test_export_daily_grid_xlsx_still_shows_day_by_day_codes(
    db, employee_a, active_status, code_present
):
    db.add(
        AttendanceEntry(employee_id=employee_a.id, date=date(2026, 9, 7), code_id=code_present.id)
    )
    db.flush()

    content = attendance_export_service.export_daily_grid_xlsx(db, 9, 2026)
    assert content[:2] == b"PK"

    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(content))
    sheet = workbook.active
    assert sheet.title == "Détail journalier"
    header = [sheet.cell(row=3, column=c).value for c in range(1, 5)]
    assert header == ["Collaborateur", "Département", 1, 2]


# ----------------------------------------------------------- monthly summary --


def test_build_monthly_summary_counts_conge_paye_and_reduces_jours_travailles(
    db, employee_a, leave_type, hr_user
):
    leave_type.code_court = "CP"
    db.flush()
    grant_balance(db, employee_a.id, leave_type.id, 2026)
    request = leave_service.create_request(
        db,
        employee_id=employee_a.id,
        leave_type_id=leave_type.id,
        date_debut=date(2026, 9, 7),
        date_fin=date(2026, 9, 8),
        commentaire=None,
        submitted_by_user_id=hr_user.id,
    )
    leave_service.approve_request(db, request, decided_by_user_id=hr_user.id, comment=None)

    rows = attendance_export_service.build_monthly_summary(db, 9, 2026)
    row = next(r for r in rows if r["employee_id"] == employee_a.id)

    jours_demandes = leave_service.jours_ouvres(db, date(2026, 9, 7), date(2026, 9, 8))
    assert row["conge_paye"] == jours_demandes
    assert row["jours_non_travailles"] == jours_demandes
    assert row["jours_travailles"] == row["jours_ouvres_mois"] - jours_demandes


def test_build_monthly_summary_counts_mission_days_without_reducing_jours_travailles(
    db, employee_a
):
    mission_code = AttendanceCode(libelle="Mission", code_court="M", couleur="#1E88E5")
    db.add(mission_code)
    db.flush()
    db.add(
        AttendanceEntry(employee_id=employee_a.id, date=date(2026, 9, 7), code_id=mission_code.id)
    )
    db.flush()

    rows = attendance_export_service.build_monthly_summary(db, 9, 2026)
    row = next(r for r in rows if r["employee_id"] == employee_a.id)

    assert row["mission"] == 1
    # Mission is worked (off-site), not an absence — doesn't reduce jours_travailles.
    assert row["jours_travailles"] == row["jours_ouvres_mois"]
    assert row["jours_non_travailles"] == 0


def test_build_monthly_summary_counts_absence_via_compte_absence_flag(db, employee_a):
    absence_code = AttendanceCode(
        libelle="Absence non justifiée", code_court="A", couleur="#E53935", compte_absence=True
    )
    db.add(absence_code)
    db.flush()
    db.add(
        AttendanceEntry(employee_id=employee_a.id, date=date(2026, 9, 7), code_id=absence_code.id)
    )
    db.flush()

    rows = attendance_export_service.build_monthly_summary(db, 9, 2026)
    row = next(r for r in rows if r["employee_id"] == employee_a.id)

    assert row["absence"] == 1
    assert row["jours_travailles"] == row["jours_ouvres_mois"] - 1


def test_export_xlsx_header_and_employee_row_match_etat_presence_template(db, employee_a):
    import openpyxl

    employee_a.matricule = "M-042"
    db.flush()

    content = attendance_export_service.export_monthly_state_xlsx(db, 9, 2026)
    workbook = openpyxl.load_workbook(BytesIO(content))
    sheet = workbook.active

    header = [sheet.cell(row=3, column=c).value for c in range(1, 34)]
    assert header == attendance_export_service._HEADERS
    # MAT / Nom / Pénom columns.
    row_values = [sheet.cell(row=4, column=c).value for c in (1, 2, 3)]
    assert row_values == ["M-042", employee_a.nom, employee_a.prenom]
