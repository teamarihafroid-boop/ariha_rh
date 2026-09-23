"""Bulk employee import: HR downloads a fixed-column .xlsx template, fills
it in outside the app (in Excel, one row per collaborateur), and uploads it
back. Unlike the pointeuse import (attendance_service), there's no
column-mapping step — the template's headers are fixed and known — so this
is a simpler upload -> preview -> confirm flow, reusing attendance_service's
generic file-parsing/temp-storage helpers rather than duplicating them.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.models import Department, Employee, EmployeeStatus, Position
from app.schemas.employee import EmployeeCreate
from app.services import employee_service
from app.services.attendance_service import _normalize  # generic accent/case-insensitive match

# Column order both the template and the parser agree on — keep in sync.
TEMPLATE_HEADERS = [
    "Matricule",
    "Nom",
    "Prénom",
    "CIN",
    "Date de naissance",
    "Lieu de naissance",
    "Téléphone",
    "Email",
    "Ville",
    "Adresse",
    "Département",
    "Poste",
    "Statut",
    "Type de contrat",
    "Catégorie",
    "CNSS",
    "Salaire de base",
    "Salaire net",
    "Date d'embauche",
    "Équipe",
]

_EXAMPLE_ROW = [
    "M-101",
    "Alami",
    "Sara",
    "BE123456",
    "15/03/1990",
    "Casablanca",
    "0600112233",
    "sara.alami@example.com",
    "Casablanca",
    "12 Rue Exemple",
    "RH",
    "Responsable RH",
    "Actif",
    "CDI",
    "cadre",
    "1234567",
    "8000",
    "10000",
    "01/06/2023",
    "",
]


@dataclass
class ImportRow:
    row_number: int
    data: EmployeeCreate | None
    display: dict[str, str]
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def build_template_xlsx(db: Session) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Collaborateurs"

    bold = Font(bold=True)
    header_fill = PatternFill(start_color="E6EEF9", end_color="E6EEF9", fill_type="solid")
    for col_idx, label in enumerate(TEMPLATE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=label)
        cell.font = bold
        cell.fill = header_fill
    for col_idx, value in enumerate(_EXAMPLE_ROW, start=1):
        sheet.cell(row=2, column=col_idx, value=value)
    for col_idx in range(1, len(TEMPLATE_HEADERS) + 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20

    reference = workbook.create_sheet("Départements et postes valides")
    reference.cell(row=1, column=1, value="Département").font = bold
    reference.cell(row=1, column=2, value="Poste").font = bold
    row_idx = 2
    departments = (
        db.query(Department).filter(Department.is_active.is_(True)).order_by(Department.nom).all()
    )
    for dept in departments:
        positions = (
            db.query(Position)
            .filter(Position.department_id == dept.id, Position.is_active.is_(True))
            .order_by(Position.intitule)
            .all()
        )
        if not positions:
            reference.cell(row=row_idx, column=1, value=dept.nom)
            row_idx += 1
            continue
        for pos in positions:
            reference.cell(row=row_idx, column=1, value=dept.nom)
            reference.cell(row=row_idx, column=2, value=pos.intitule)
            row_idx += 1
    reference.column_dimensions["A"].width = 28
    reference.column_dimensions["B"].width = 28

    statuses = workbook.create_sheet("Statuts valides")
    statuses.cell(row=1, column=1, value="Statut").font = bold
    for i, status in enumerate(
        db.query(EmployeeStatus).order_by(EmployeeStatus.libelle).all(), start=2
    ):
        statuses.cell(row=i, column=1, value=status.libelle)
    statuses.column_dimensions["A"].width = 20

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _parse_date(raw: str) -> tuple[date | None, str | None]:
    raw = raw.strip()
    if not raw:
        return None, None
    # Excel-native date cells come back from read_table as a Python str(datetime),
    # e.g. "1990-03-15 00:00:00" — try that before the typed "JJ/MM/AAAA" form.
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            # A birth/hire date has no timezone concept — naive is correct here.
            return datetime.strptime(raw, fmt).date(), None  # noqa: DTZ007
        except ValueError:
            continue
    return None, f"date non reconnue : {raw!r} (attendu JJ/MM/AAAA)"


def _parse_decimal(raw: str) -> tuple[Decimal | None, str | None]:
    raw = raw.strip().replace(" ", "").replace(",", ".")
    if not raw:
        return None, None
    try:
        return Decimal(raw), None
    except InvalidOperation:
        return None, f"montant non reconnu : {raw!r}"


def _build_department_lookup(db: Session) -> dict[str, int]:
    return {
        _normalize(d.nom): d.id
        for d in db.query(Department).filter(Department.is_active.is_(True)).all()
    }


def _build_position_lookup(db: Session) -> dict[str, list[tuple[int, int | None]]]:
    lookup: dict[str, list[tuple[int, int | None]]] = {}
    for p in db.query(Position).filter(Position.is_active.is_(True)).all():
        lookup.setdefault(_normalize(p.intitule), []).append((p.id, p.department_id))
    return lookup


def _build_status_lookup(db: Session) -> dict[str, int]:
    return {_normalize(s.libelle): s.id for s in db.query(EmployeeStatus).all()}


def _default_status_id(db: Session) -> int | None:
    default = db.query(EmployeeStatus).filter_by(libelle="Actif").first()
    return default.id if default else None


def _is_probable_duplicate(db: Session, nom: str, prenom: str, date_naissance: date | None) -> bool:
    query = db.query(Employee).filter(Employee.nom.ilike(nom), Employee.prenom.ilike(prenom))
    if date_naissance is not None:
        query = query.filter(Employee.date_naissance == date_naissance)
    return db.query(query.exists()).scalar()


def _cell(row: dict[str, str], key: str) -> str:
    return (row.get(key) or "").strip()


def parse_rows(db: Session, rows: list[dict[str, str]]) -> list[ImportRow]:
    dept_lookup = _build_department_lookup(db)
    position_lookup = _build_position_lookup(db)
    status_lookup = _build_status_lookup(db)
    default_status_id = _default_status_id(db)

    results: list[ImportRow] = []
    for i, row in enumerate(rows, start=2):  # row 1 is the header
        nom, prenom = _cell(row, "Nom"), _cell(row, "Prénom")
        display = {"Nom": nom, "Prénom": prenom, "Département": _cell(row, "Département")}
        errors: list[str] = []
        warnings: list[str] = []

        if not nom or not prenom:
            errors.append("Nom et Prénom sont obligatoires.")

        department_id = None
        dept_raw = _cell(row, "Département")
        if dept_raw:
            department_id = dept_lookup.get(_normalize(dept_raw))
            if department_id is None:
                errors.append(f"Département inconnu : {dept_raw!r}")

        position_id = None
        poste_raw = _cell(row, "Poste")
        if poste_raw:
            candidates = position_lookup.get(_normalize(poste_raw), [])
            if not candidates:
                errors.append(f"Poste inconnu : {poste_raw!r}")
            elif department_id is not None:
                matching = [pid for pid, did in candidates if did == department_id]
                position_id = matching[0] if matching else candidates[0][0]
                if not matching:
                    warnings.append(
                        f"Poste {poste_raw!r} n'appartient pas au département {dept_raw!r} indiqué."
                    )
            else:
                position_id = candidates[0][0]
                if len(candidates) > 1:
                    warnings.append(
                        f"Poste {poste_raw!r} ambigu (plusieurs départements) — premier utilisé."
                    )

        status_raw = _cell(row, "Statut")
        if status_raw:
            status_id = status_lookup.get(_normalize(status_raw))
            if status_id is None:
                errors.append(f"Statut inconnu : {status_raw!r}")
        else:
            status_id = default_status_id

        date_naissance, date_naissance_err = _parse_date(_cell(row, "Date de naissance"))
        if date_naissance_err:
            errors.append(date_naissance_err)
        date_embauche, date_embauche_err = _parse_date(_cell(row, "Date d'embauche"))
        if date_embauche_err:
            errors.append(date_embauche_err)

        salaire_base, salaire_base_err = _parse_decimal(_cell(row, "Salaire de base"))
        if salaire_base_err:
            errors.append(salaire_base_err)
        salaire_net, salaire_net_err = _parse_decimal(_cell(row, "Salaire net"))
        if salaire_net_err:
            errors.append(salaire_net_err)

        if not errors and _is_probable_duplicate(db, nom, prenom, date_naissance):
            warnings.append("Un collaborateur avec le même nom (et date de naissance) existe déjà.")

        data = None
        if not errors:
            data = EmployeeCreate(
                matricule=_cell(row, "Matricule") or None,
                nom=nom,
                prenom=prenom,
                cin=_cell(row, "CIN") or None,
                date_naissance=date_naissance,
                lieu_naissance=_cell(row, "Lieu de naissance") or None,
                telephone=_cell(row, "Téléphone") or None,
                email=_cell(row, "Email") or None,
                ville=_cell(row, "Ville") or None,
                adresse=_cell(row, "Adresse") or None,
                department_id=department_id,
                position_id=position_id,
                status_id=status_id,
                type_contrat=_cell(row, "Type de contrat") or None,
                categorie_professionnelle=_cell(row, "Catégorie") or None,
                cnss=_cell(row, "CNSS") or None,
                salaire_base=salaire_base,
                salaire_net=salaire_net,
                date_embauche=date_embauche,
                equipe=_cell(row, "Équipe") or None,
            )

        results.append(
            ImportRow(row_number=i, data=data, display=display, errors=errors, warnings=warnings)
        )
    return results


@dataclass
class ImportOutcome:
    created: int
    skipped: list[dict]


def create_employees(db: Session, rows: list[ImportRow]) -> ImportOutcome:
    created = 0
    skipped: list[dict] = []
    for row in rows:
        if row.data is None:
            skipped.append(
                {
                    "row_number": row.row_number,
                    "display": row.display,
                    "reason": "; ".join(row.errors),
                }
            )
            continue
        try:
            employee_service.create_employee(db, row.data)
            created += 1
        except employee_service.EmployeeServiceError as exc:
            skipped.append(
                {"row_number": row.row_number, "display": row.display, "reason": str(exc)}
            )
    return ImportOutcome(created=created, skipped=skipped)
